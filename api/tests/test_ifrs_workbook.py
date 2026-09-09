from datetime import datetime
from decimal import Decimal

import openpyxl

from screener.audit import _source_values
from screener.normalize import build_snapshot
from screener.sources.ifrs_workbook import build_adidas_companyfacts
from screener.sync import _source


def _statement(sheet, report_year, prior_year, rows):
    sheet.append(["adidas AG statement"])
    sheet.append([])
    sheet.append([])
    sheet.append([])
    sheet.append([None, "Note", f"Year ending Dec. 31, {report_year}",
                  f"Year ending Dec. 31, {prior_year}"])
    for label, current, prior in rows:
        sheet.append([label, None, current, prior])


def _report(path, report_year, prior_year, *, current_revenue, prior_revenue):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    workbook.properties.modified = datetime(report_year + 1, 3, 1)
    balance = workbook.create_sheet("cfs-cons-stat-fin-position")
    _statement(balance, report_year, prior_year, [
        ("Cash and cash equivalents", 100, 90),
        ("Short-term financial assets", "–", "–"),
        ("Accounts receivable", 250, 220),
        ("Inventories", 300, 280),
        ("Total current assets", 900, 850),
        ("Property, plant, and equipment", 500, 480),
        ("Goodwill", 100, 95),
        ("Other intangible assets", 50, 45),
        ("Long-term financial assets", 25, 20),
        ("Total assets", 2000, 1900),
        ("Short-term borrowings", 40, 35),
        ("Current lease liabilities", 20, 18),
        ("Total current liabilities", 600, 570),
        ("Long-term borrowings", 300, 320),
        ("Non-current lease liabilities", 80, 75),
        ("Total non-current liabilities", 500, 510),
        ("Shareholders’ equity", 850, 790),
        ("Non-controlling interests", 50, 30),
        ("Total equity", 900, 820),
        ("Total liabilities and equity", 2000, 1900),
    ])
    income = workbook.create_sheet("cfs-cons-stat-income")
    _statement(income, report_year, prior_year, [
        ("Net sales", current_revenue, prior_revenue),
        ("Gross profit", 500, 450),
        ("Operating profit", 180, 160),
        ("Income before taxes", 150, 140),
        ("Income taxes", 35, 32),
        ("Net income", 115, 108),
        ("Net income attributable to shareholders", 110, 104),
        ("Net income attributable to non-controlling interests", 5, 4),
        ("Basic earnings per share from continuing and discontinued operations (in €)",
         1.10, 1.04),
        ("Diluted earnings per share from continuing and discontinued operations (in €)",
         1.05, 1.00),
    ])
    cash = workbook.create_sheet("cfs-cons-stat-cash-flows")
    _statement(cash, report_year, prior_year, [
        ("Depreciation, amortization, and impairment losses", 70, 65),
        ("Interest expense", 15, 14),
        ("Change in receivables and other assets", -20, -15),
        ("Change in inventories", -10, -8),
        ("Change in accounts payable and other liabilities", 5, 7),
        ("Income taxes paid", -30, -28),
        ("Cash flows from operating activities", 160, 150),
        ("Purchase of other intangible assets", -10, -9),
        ("Purchase of property, plant, and equipment", -60, -55),
        ("Dividend paid to shareholders of adidas AG", -25, -20),
    ])
    workbook.save(path)


def test_adidas_adapter_preserves_eur_rows_and_latest_comparative(tmp_path):
    _report(tmp_path / "cons-financial-statements-adidas-ar24.xlsx", 2024, 2023,
            current_revenue=1200, prior_revenue=1100)
    _report(tmp_path / "cons-financial-statements-adidas-ar25.xlsx", 2025, 2024,
            current_revenue=1300, prior_revenue=1250)

    companyfacts = build_adidas_companyfacts(tmp_path)
    snapshot = build_snapshot("ADS.DE", "IFRS-ADIDAS-AG", companyfacts)

    assert snapshot.reporting_currency == "EUR"
    assert snapshot.annual_revenue[2024].value == 1_250_000_000
    assert snapshot.annual_revenue[2024].provenance.document.endswith("ar25.xlsx")
    assert snapshot.annual_revenue[2024].provenance.tag == "adidas-workbook:Net sales"
    assert snapshot.annual_revenue[2024].provenance.unit == "EUR"
    assert snapshot.total_assets.value == 2_000_000_000
    assert snapshot.annual_eps[2025].value == Decimal("1.05")

    source = _source(snapshot.total_assets)
    assert source["canonical_tag"] == "Assets"
    assert source["document"].endswith("ar25.xlsx")
    assert _source_values(companyfacts, source) == [2_000_000_000.0]

    # These consolidated statement downloads carry EPS but no weighted diluted
    # count. Total cash evidence still renders; per-share owner figures stay missing.
    annual = snapshot.owner_earnings.annual[2025]
    assert annual.diluted_shares is None
    assert annual.free_cash_flow.value == 100_000_000
    assert annual.free_cash_flow_per_share is None

    nil = companyfacts["facts"]["canonical"]["ShortTermInvestments"]["units"]["EUR"]
    assert any(entry["val"] == 0 and entry["_source_nil_marker"] for entry in nil)


def test_adidas_cash_payment_signs_follow_canonical_magnitude(tmp_path):
    _report(tmp_path / "cons-financial-statements-adidas-ar25.xlsx", 2025, 2024,
            current_revenue=1300, prior_revenue=1250)
    facts = build_adidas_companyfacts(tmp_path)["facts"]["canonical"]

    for tag in (
        "PaymentsToAcquirePropertyPlantAndEquipment",
        "PaymentsToAcquireIntangibleAssets",
        "PaymentsOfDividendsCommonStock",
        "IncomeTaxesPaidNet",
    ):
        assert all(entry["val"] >= 0 for entry in facts[tag]["units"]["EUR"])
        assert all("sign normalized" in entry["_source_concept"]
                   for entry in facts[tag]["units"]["EUR"])
